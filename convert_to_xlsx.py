#!/usr/bin/env python3
"""
Convert BMW scraper JSON to XLSX
"""

import json
import sys
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("Missing dependency. Install with:")
    print("  pip install openpyxl")
    sys.exit(1)


def json_to_xlsx(json_file, xlsx_file, enriched_csv=None):
    """Convert JSON to formatted XLSX with optional RDW data"""

    # Load JSON
    with open(json_file, 'r', encoding='utf-8') as f:
        vehicles = json.load(f)

    print(f"Loaded {len(vehicles)} vehicles from {json_file}")

    # Load RDW data if provided
    rdw_data = {}
    if enriched_csv and Path(enriched_csv).exists():
        import csv
        with open(enriched_csv, 'r', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row in reader:
                vehicle_id = row.get('vehicle_id', '')
                rdw_data[vehicle_id] = {
                    'catalogue_price': row.get('catalogue_price', ''),
                    'discount_percent': row.get('discount_percent', ''),
                    'rdw_model': row.get('rdw_model', ''),
                    'first_registration': row.get('first_registration', '')
                }
        print(f"Loaded RDW data for {len(rdw_data)} vehicles from {enriched_csv}")

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "BMW i5 Listings"

    # Define headers
    headers = [
        'License Plate', 'Vehicle ID', 'URL', 'Name', 'Model', 'Chassis',
        'Price', 'Catalogue Price', 'Discount %', 'Mileage', 'Year',
        'Fuel', 'Transmission', 'Engine', 'Color', 'Dealer', 'City', 'First Registration'
    ]

    # Write headers with formatting
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write data
    for row_idx, vehicle in enumerate(vehicles, 2):
        vehicle_id = vehicle.get('vehicleId', '')
        rdw = rdw_data.get(vehicle_id, {})

        ws.cell(row=row_idx, column=1, value=vehicle.get('license_plate', ''))
        ws.cell(row=row_idx, column=2, value=vehicle_id)

        # URL as hyperlink
        url = f"https://occasions.bmw.nl/bmw/zoeken/resultaten/details/id/{vehicle_id}" if vehicle_id else ''
        ws.cell(row=row_idx, column=3, value=url)
        if url:
            from openpyxl.styles import Font as ExcelFont
            ws.cell(row=row_idx, column=3).hyperlink = url
            ws.cell(row=row_idx, column=3).font = ExcelFont(color="0563C1", underline="single")

        ws.cell(row=row_idx, column=4, value=vehicle.get('name', ''))
        ws.cell(row=row_idx, column=5, value=vehicle.get('model', ''))
        ws.cell(row=row_idx, column=6, value=vehicle.get('chassis', ''))

        # Price as number
        price = vehicle.get('price', '')
        if price:
            ws.cell(row=row_idx, column=7, value=int(price))
            ws.cell(row=row_idx, column=7).number_format = '€#,##0'

        # Catalogue price as number
        cat_price = rdw.get('catalogue_price', '')
        if cat_price:
            ws.cell(row=row_idx, column=8, value=int(cat_price))
            ws.cell(row=row_idx, column=8).number_format = '€#,##0'

        # Discount percentage
        discount = rdw.get('discount_percent', '')
        if discount:
            ws.cell(row=row_idx, column=9, value=float(discount))
            ws.cell(row=row_idx, column=9).number_format = '0.0"%"'

        # Mileage as number
        mileage = vehicle.get('mileage', '')
        if mileage:
            ws.cell(row=row_idx, column=10, value=int(mileage))
            ws.cell(row=row_idx, column=10).number_format = '#,##0'

        # Year
        year = vehicle.get('datePartOne', '')[:4] if vehicle.get('datePartOne') else ''
        ws.cell(row=row_idx, column=11, value=year)

        ws.cell(row=row_idx, column=12, value=vehicle.get('fuel', ''))
        ws.cell(row=row_idx, column=13, value=vehicle.get('transmission', ''))
        ws.cell(row=row_idx, column=14, value=vehicle.get('engine', ''))
        ws.cell(row=row_idx, column=15, value=vehicle.get('color', ''))
        ws.cell(row=row_idx, column=16, value=vehicle.get('dealerName', ''))
        ws.cell(row=row_idx, column=17, value=vehicle.get('dealerCity', ''))

        # First registration
        first_reg = rdw.get('first_registration', '')
        if first_reg and len(first_reg) == 8:
            formatted_date = f"{first_reg[6:8]}-{first_reg[4:6]}-{first_reg[0:4]}"
            ws.cell(row=row_idx, column=18, value=formatted_date)

    # Auto-size columns
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Save
    wb.save(xlsx_file)
    print(f"Saved to: {xlsx_file}")
    print(f"  Rows: {len(vehicles)} vehicles")
    print(f"  Columns: {len(headers)}")


def main():
    # Find most recent raw JSON
    json_files = list(Path('.').glob('vehicles_raw_*.json'))

    if not json_files:
        print("No vehicles_raw_*.json files found")
        return 1

    # Use most recent
    json_file = sorted(json_files)[-1]

    # Find most recent enriched CSV
    enriched_csv = None
    enriched_files = list(Path('.').glob('vehicles_enriched_*.csv'))
    if enriched_files:
        enriched_csv = sorted(enriched_files)[-1]

    # Generate output filename
    date_str = datetime.now().strftime('%Y-%m-%d_%H%M%S')
    xlsx_file = f"vehicles_{date_str}.xlsx"

    json_to_xlsx(json_file, xlsx_file, enriched_csv)

    return 0


if __name__ == '__main__':
    sys.exit(main())
